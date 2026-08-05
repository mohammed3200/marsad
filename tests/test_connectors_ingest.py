import datetime
import json
import tempfile
import unittest
from pathlib import Path

import connectors
from connectors import guess_dept, is_internal_file, read_file_to_report

ARABIC = "أنجز فريق الموقع ٨٠٪ من أعمال الحفر اليوم"


class GuessDeptTests(unittest.TestCase):
    def test_english_keywords(self):
        self.assertEqual(guess_dept("ran_daily.docx"), "ran")
        self.assertEqual(guess_dept("core_network.csv"), "core")
        self.assertEqual(guess_dept("quality_sites.xlsx"), "quality")
        self.assertEqual(guess_dept("safety_site.pdf"), "safety")

    def test_ran_does_not_match_inside_ordinary_words(self):
        for name in ("random_stuff.txt", "transfer_log.txt",
                     "grant_letter.txt", "France_visit.txt"):
            with self.subTest(name=name):
                self.assertEqual(guess_dept(name), "admin")

    def test_normal_arabic_spellings_are_recognised(self):
        self.assertEqual(guess_dept("تكاليف_المشروع.xlsx"), "cost")
        self.assertEqual(guess_dept("إنشاء_الأبراج.docx"), "civil")
        self.assertEqual(guess_dept("تقرير_الجودة.pdf"), "quality")
        self.assertEqual(guess_dept("السلامة_اليومي.txt"), "safety")

    def test_unknown_falls_back_to_admin(self):
        self.assertEqual(guess_dept("notes.txt"), "admin")

    def test_arabic_keywords_do_not_match_as_mere_substrings(self):
        # Regression for the عقد collision: it should not match in unrelated words
        # معقد (muʿaqqad) = complicated, not a contract
        self.assertEqual(guess_dept("تقرير_معقد.pdf"), "admin")
        # العقد here means "decade", not singular contract — a genuine homograph
        # (ال + عقد is identical either way), resolved by matching عقد/عقود as
        # bare tokens only, never through proclitic-stripped "ال" forms.
        self.assertEqual(guess_dept("خطة_العقد_القادم.xlsx"), "admin")
        # عقدة (ʿuqda) = psychological complex, not a contract
        self.assertEqual(guess_dept("عقدة_نفسية.docx"), "admin")

    def test_latin_keywords_match_when_attached_to_digits(self):
        # Regression for digit handling: RAN2024 should still match "ran"
        self.assertEqual(guess_dept("RAN2024_report.docx"), "ran")
        # core5g should match "core" (and would have matched via "network" too)
        self.assertEqual(guess_dept("core5g_network.csv"), "core")
        # 5Gcore should not match (5G prefix, digit-glued to "core", creates "gcore" token)
        # This is a known gap: false negative (manual triage) not false positive
        self.assertEqual(guess_dept("5Gcore_report.csv"), "admin")

    def test_quality_keyword_does_not_match_existing_data(self):
        # Regression for جودة collision with موجودة (existing)
        # بيانات_موجودة = "existing data"
        self.assertEqual(guess_dept("بيانات_موجودة.xlsx"), "admin")
        # Positive cases: bare جودة, and الجودة via proclitic stripping of "ال"
        self.assertEqual(guess_dept("تقرير_الجودة.pdf"), "quality")
        self.assertEqual(guess_dept("جودة_المشروع.pdf"), "quality")
        self.assertEqual(guess_dept("مراقبة_جودة_الخرسانة.pdf"), "quality")

    def test_cost_keyword_does_not_match_northern_or_probability(self):
        # Regression for مالية collision with شمالية (northern), احتمالية (probability),
        # and عمالية (labor) — all common words that would misroute files
        # المنطقة_الشمالية = "Northern region" (very common in telecom regional rollout context)
        self.assertEqual(guess_dept("المنطقة_الشمالية.pdf"), "admin")
        # تحليل_احتمالية_المخاطر = "risk probability analysis"
        self.assertEqual(guess_dept("تحليل_احتمالية_المخاطر.xlsx"), "admin")
        # نزاع_عمالية = "labor dispute"
        self.assertEqual(guess_dept("نزاع_عمالية.pdf"), "admin")
        # Positive case: genuine cost reports still route via more specific keywords
        self.assertEqual(guess_dept("تقرير_التكاليف.xlsx"), "cost")
        self.assertEqual(guess_dept("الميزانية_السنوية.pdf"), "cost")
        # Restored مالية: matches بare and via "ال"-stripping, without the collisions above
        self.assertEqual(guess_dept("الشؤون_المالية.pdf"), "cost")
        self.assertEqual(guess_dept("الموارد_المالية_للمشروع.xlsx"), "cost")

    def test_token_matching_closes_the_whole_collision_class(self):
        # محضر الجلسة المعقودة = "minutes of the convened meeting" — معقودة is a
        # different word from عقد (passive participle of عقد "to convene/hold"),
        # not an affixed form of it, so token equality rejects it.
        self.assertEqual(guess_dept("محضر_الجلسة_المعقودة.pdf"), "admin")
        # اجتماع غير معقود = "meeting not held"
        self.assertEqual(guess_dept("اجتماع_غير_معقود.pdf"), "admin")
        # حبل مجدول = "twisted/braided rope" (reinforcement wire) — مجدول is a
        # different word from جدول, not جدول with a proclitic prefix.
        self.assertEqual(guess_dept("حبل_مجدول.pdf"), "admin")
        self.assertEqual(guess_dept("خيط_مجدول_للتسليح.docx"), "admin")

    def test_contract_keyword_restored_and_safe(self):
        # عقود (plural) — unambiguous, matches directly.
        self.assertEqual(guess_dept("عقود_المقاولين.pdf"), "contract")
        # Singular عقد as a bare token (no definite article) — the ordinary way a
        # contract filename is named — routes correctly.
        self.assertEqual(guess_dept("عقد_المقاول_الرئيسي.pdf"), "contract")

    def test_safety_keyword_affixed_forms(self):
        self.assertEqual(guess_dept("السلامة_اليومي.txt"), "safety")
        # بالسلامة ("بال" proclitic cluster + سلامة) — "commitment to safety"
        self.assertEqual(guess_dept("الالتزام_بالسلامة.pdf"), "safety")

    def test_civil_keyword_spelling_variants(self):
        self.assertEqual(guess_dept("إنشاء_الأبراج.docx"), "civil")
        self.assertEqual(guess_dept("انشاء_الأبراج.docx"), "civil")

    def test_schedule_keyword_with_definite_article(self):
        self.assertEqual(guess_dept("الجدول_الزمني.xlsx"), "schedule")


class IngestTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_every_supported_format_round_trips(self):
        import docx, openpyxl
        from reportlab.pdfgen import canvas

        d = docx.Document(); d.add_paragraph(ARABIC); d.save(self.dir / "ran_daily.docx")
        wb = openpyxl.Workbook(); wb.active.append(["البند", ARABIC])
        wb.save(self.dir / "quality_sites.xlsx")
        (self.dir / "core_network.csv").write_text(
            f"القسم,الحالة\nالنواة,{ARABIC}\n", encoding="utf-8")
        (self.dir / "تقرير_الجودة.txt").write_text(ARABIC, encoding="utf-8")
        (self.dir / "plain.json").write_text(
            json.dumps({"ملاحظة": ARABIC}, ensure_ascii=False), encoding="utf-8")
        c = canvas.Canvas(str(self.dir / "safety_site.pdf"))
        c.drawString(72, 720, "RAN site safety inspection completed")
        c.save()

        today = datetime.date.today().isoformat()
        for name, dept in (("ran_daily.docx", "ran"),
                           ("quality_sites.xlsx", "quality"),
                           ("core_network.csv", "core"),
                           ("تقرير_الجودة.txt", "quality"),
                           ("plain.json", "admin"),
                           ("safety_site.pdf", "safety")):
            with self.subTest(name=name):
                rep = read_file_to_report(self.dir / name)
                self.assertIsNotNone(rep)
                self.assertEqual(set(rep),
                                 {"id", "source", "dept", "from", "date", "content"})
                self.assertEqual(rep["source"], "upload")
                self.assertEqual(rep["date"], today)
                self.assertEqual(rep["dept"], dept)

    def test_empty_file_is_rejected(self):
        (self.dir / "empty.txt").write_text("", encoding="utf-8")
        self.assertIsNone(read_file_to_report(self.dir / "empty.txt"))

    def test_a_corrupt_file_is_rejected_not_narrated(self):
        """An unreadable file must not become a report.

        The readers used to return their error message as the content — a
        truthy Arabic string that sailed past the `if not content` guard. The
        file then entered the queue and was passed verbatim to the AI engine
        by _format_reports, so the model was asked to analyse
        «[خطأ في قراءة Excel: الملف تالف…]» as if it were a field report.
        """
        for name, blob in (("broken.xlsx", b"not a zip file at all"),
                           ("broken.docx", b"neither is this"),
                           ("broken.pdf",  b"%PDF-1.4 truncated")):
            with self.subTest(name=name):
                p = self.dir / name
                p.write_bytes(blob)
                self.assertIsNone(read_file_to_report(p))

    def test_a_legacy_xls_is_not_advertised(self):
        """openpyxl cannot read the legacy BIFF format, so claiming .xls in
        DOC_PATTERNS and the file picker promised something that always
        failed. Better to not offer it than to offer it broken."""
        import connectors
        self.assertNotIn("*.xls", connectors.DOC_PATTERNS)
        self.assertIn("*.xlsx", connectors.DOC_PATTERNS)

    def test_internal_files_can_never_be_ingested(self):
        from core.paths import BUNDLE_DIR, DATA_DIR
        for path in (DATA_DIR / "settings.json",
                     BUNDLE_DIR / "settings.example.json",
                     DATA_DIR / "data" / "contacts.json",
                     DATA_DIR / "reports" / "latest.json"):
            with self.subTest(path=path.name):
                self.assertTrue(is_internal_file(path))
                self.assertIsNone(read_file_to_report(path))


# ── moved here when the web API was removed ────────────────────────────
# These test connectors.py, not the API. They lived in
# tests/test_api_uploads.py only because that file was written when the
# upload endpoint was the thing under test; deleting it wholesale would
# have silently dropped real coverage of the read caps.


class ReadCapTests(unittest.TestCase):
    """read_file_to_report writes no user state, so it needs no isolation."""

    def _huge(self, name, text):
        import tempfile
        from pathlib import Path
        p = Path(tempfile.mkdtemp()) / name
        p.write_text(text, encoding="utf-8")
        return p

    def test_txt_content_is_capped(self):
        import connectors
        p = self._huge("ran_huge.txt", "ب" * (connectors.MAX_CHARS + 5000))
        rep = connectors.read_file_to_report(p)
        self.assertLessEqual(len(rep["content"]), connectors.MAX_CHARS + 200)
        self.assertIn("اقتُطع", rep["content"])

    def test_content_under_the_cap_is_untouched(self):
        import connectors
        body = "تقرير قصير عن الموقع"
        p = self._huge("ran_small.txt", body)
        rep = connectors.read_file_to_report(p)
        self.assertEqual(rep["content"], body)
        self.assertNotIn("اقتُطع", rep["content"])

    def test_cap_helper_boundary(self):
        import connectors
        exact = "ب" * connectors.MAX_CHARS
        self.assertEqual(connectors._cap(exact), exact)
        over = "ب" * (connectors.MAX_CHARS + 1)
        capped = connectors._cap(over)
        self.assertTrue(capped.startswith("ب" * connectors.MAX_CHARS))
        self.assertIn("اقتُطع", capped)

    def test_cap_helper_tolerates_empty_and_none(self):
        import connectors
        self.assertEqual(connectors._cap(""), "")
        self.assertEqual(connectors._cap(None), "")

    def test_docx_content_is_capped(self):
        import tempfile
        from pathlib import Path
        import connectors
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        p = Path(tempfile.mkdtemp()) / "ran_huge.docx"
        d = docx.Document()
        for _ in range(60):
            d.add_paragraph("ب" * 1000)
        d.save(p)
        rep = connectors.read_file_to_report(p)
        self.assertLessEqual(len(rep["content"]), connectors.MAX_CHARS + 200)
        self.assertIn("اقتُطع", rep["content"])

if __name__ == "__main__":
    unittest.main()

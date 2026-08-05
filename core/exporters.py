"""
marsad report exporters — PDF + Excel.

UI-agnostic standalone functions taking (results, output_path). Moved out of the
root package __init__ so the old root package can be removed. Both read the same
`results` dict (keyed by agent id) that the dashboard and connectors use.
"""
import os, re, datetime
from pathlib import Path

from core.status import tier


def _num(value, default=0.0) -> float:
    """First number in a value the model may have returned as '80%' or '4.2 مليون'.

    A blunt heuristic: '2024-2025' also parses as 2024.0. Fine for the
    percentage/count fields this is used on today; do not point it at
    currency or date-range fields without re-checking that assumption."""
    if isinstance(value, bool):
        return float(default)
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        m = re.search(r"-?\d+(?:\.\d+)?", value.replace(",", ""))
        if m:
            return float(m.group())
    return float(default)


def _int(value, default=0) -> int:
    return int(_num(value, default))


def _pct(value) -> str:
    """Format a percentage without ever producing '80%%' or 'None%'."""
    return f"{_num(value, 0.0):g}%"


def _rows(value) -> list:
    return value if isinstance(value, list) else []


def _obj(value) -> dict:
    return value if isinstance(value, dict) else {}


def _cell(value):
    """A value safe to assign to an openpyxl cell.

    openpyxl accepts str/int/float/bool/None/datetime but raises on a raw
    list or dict — reachable here because tier() and _obj()/_rows() only
    guard the *classification* of a model-returned field (health/level/
    status), not the raw value written into the sheet next to it (e.g. a
    risk `level` of `["عالية"]`). Anything else is stringified, matching
    the str(v) already used a few rows down for the schedule/finance sheets.
    """
    return value if isinstance(value, (str, int, float, bool, type(None))) else str(value)


def _wrap_rtl(text: str, measure, width: float) -> list:
    """Break `text` into lines that fit `width`, working on the LOGICAL text.

    `measure(str) -> float` reports the rendered width of a candidate line.

    Kept module-level and dependency-free so it can be tested without
    reportlab, fonts, or a PDF: the ordering bug this exists to prevent is
    pure list logic, and the reshaping around it is not what got it wrong.
    """
    lines, cur = [], ""
    for word in str(text or "").split():
        trial = f"{cur} {word}".strip()
        if not cur or measure(trial) <= width:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def export_pdf(results: dict, output_path: str) -> str:
    """Branded Arabic PDF — Light Executive Report identity (white paper,
    hairlines, one teal accent, status as colored words/dots).

    reportlab ships no Arabic glyphs and no shaping: we register the bundled
    Noto fonts and run Arabic through arabic_reshaper + python-bidi (`ar()`).
    """
    from xml.sax.saxutils import escape as _esc

    from reportlab.lib.pagesizes   import A4
    from reportlab.lib             import colors
    from reportlab.lib.units       import cm
    from reportlab.platypus        import (SimpleDocTemplate, Paragraph, Spacer,
                                           Table, TableStyle, HRFlowable)
    from reportlab.lib.styles      import ParagraphStyle
    from reportlab.lib.enums       import TA_RIGHT, TA_CENTER, TA_LEFT
    from reportlab.pdfbase         import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.graphics.shapes import Drawing, Circle

    import arabic_reshaper
    from bidi.algorithm import get_display

    from .paths import bundle_path

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)

    # ── bundled fonts (the base-14 fonts have no Arabic at all; Noto Kufi is
    #    intentionally unused — reportlab renders nothing from its TTFs) ──
    for name, fn in (("Naskh",      "NotoSansArabic-Regular.ttf"),
                     ("Naskh-Bold", "NotoSansArabic-Bold.ttf"),
                     ("Mono",       "JetBrainsMono-Regular.ttf")):
        pdfmetrics.registerFont(TTFont(name, str(bundle_path("assets", "fonts", fn))))

    def _has_ar(s):
        return any("؀" <= ch <= "ۿ" for ch in s)

    def ar(text, width=None, font="Naskh", size=11):
        """Arabic-ready string for a reportlab Paragraph: reshape → bidi → escape.

        Escaping must come last: XML entities contain Latin letters and ASCII
        punctuation, so reordering them inside an RTL run emits ';pma&'.

        En/em dashes become hyphen-minus first: U+2013/2014 break the BiDi
        number run (renders "8060" out of "60–80%"), while ES separators
        (-, /, .) keep digits as one LTR run.

        `width` is REQUIRED for any text that can wrap to more than one line.
        get_display() reverses the whole string, so if reportlab then breaks
        that already-reversed run into lines, the lines come out bottom-up —
        the last sentence renders first. Passing the column width makes this
        wrap the *logical* text itself and reverse each line separately, which
        keeps line order. Found on the first real model run: every multi-line
        block in the PDF — the executive summary and every risk-table cell —
        read bottom-to-top. Short fixtures never wrapped, so no test saw it.
        """
        s = str(text if text is not None else "")
        s = s.replace("–", "-").replace("—", "-")
        if not _has_ar(s):
            return _esc(s)
        if width is None:
            return _esc(get_display(arabic_reshaper.reshape(s)))

        def measure(candidate):
            return pdfmetrics.stringWidth(
                get_display(arabic_reshaper.reshape(candidate)), font, size)

        return "<br/>".join(
            _esc(get_display(arabic_reshaper.reshape(line)))
            for line in _wrap_rtl(s, measure, width))

    # ── identity tokens (backend/theme.py) ──
    INK   = colors.HexColor("#141A22")
    INK2  = colors.HexColor("#5A6675")
    INK3  = colors.HexColor("#8A94A1")
    ACC   = colors.HexColor("#0E6E60")
    LINE  = colors.HexColor("#E4E8EC")
    LINEH = colors.HexColor("#CFD6DD")
    GOOD  = colors.HexColor("#1E7A52")
    WARN  = colors.HexColor("#946200")
    BAD   = colors.HexColor("#B4232A")

    _PDF_TIER = {"good": GOOD, "warn": WARN, "bad": BAD, "neutral": INK2}

    def status_color(lit):
        return _PDF_TIER[tier(lit)]

    CONTENT_W = 17 * cm

    def PS(name, font="Naskh", size=10, color=INK, align=TA_RIGHT, leading=None, **kw):
        return ParagraphStyle(name, fontName=font, fontSize=size, textColor=color,
                              alignment=align, leading=leading or size * 1.6, **kw)

    def val(text, align=TA_CENTER, color=INK, size=10, bold=False):
        """Value cell — JetBrains Mono for Latin/numbers (has ↑↓→), Naskh for Arabic."""
        s = str(text if text is not None else "")
        font = ("Naskh-Bold" if bold else "Naskh") if _has_ar(s) else "Mono"
        return Paragraph(ar(s), PS(f"v{font}{size}", font, size, color, align))

    def dot(color, d=0.22 * cm):
        dr = Drawing(d, d)
        dr.add(Circle(d / 2, d / 2, d / 2, fillColor=color, strokeColor=None))
        return dr

    def section(title):
        """Bold section title with a short accent rule directly under it (right)."""
        return [Spacer(1, 10),
                Paragraph(ar(title), PS("sec", "Naskh-Bold", 13, INK)),
                HRFlowable(width=2.2 * cm, thickness=2, color=ACC, hAlign="RIGHT")]

    def rows_style(n, header=False):
        cmds = [("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2)]
        if header:
            cmds.append(("LINEBELOW", (0, 0), (-1, 0), 0.7, LINEH))
        for i in range(1 if header else 0, n):
            cmds.append(("LINEBELOW", (0, i), (-1, i), 0.4, LINE))
        return cmds

    AR_D = "٠١٢٣٤٥٦٧٨٩"
    def arabic_num(n):
        return "".join(AR_D[int(d)] if d.isdigit() else d for d in str(n))

    chief = _obj(results.get("chief"))
    risk  = _obj(results.get("risk"))
    sched = _obj(results.get("schedule"))
    fin   = _obj(results.get("cost"))
    health = chief.get("overall_health", "غير محدد")

    story = []

    # ── title block ──
    story.append(Table([[
        Paragraph(ar("تقرير حالة المشروع"), PS("t", "Naskh-Bold", 20, INK)),
        Paragraph(datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
                  PS("dt", "Mono", 9, INK3, TA_LEFT))]],
        colWidths=[CONTENT_W - 4 * cm, 4 * cm],
        style=TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                          ("LEFTPADDING", (0, 0), (-1, -1), 0),
                          ("RIGHTPADDING", (0, 0), (-1, -1), 0)])))
    story.append(Paragraph(ar("مرصد — ذكاء المشاريع · LTT 4G/5G"), PS("br", "Naskh", 9, INK3)))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=2, color=ACC))

    # ── health line ──
    story.append(Spacer(1, 12))
    story.append(Table([[
        Paragraph(ar("الحالة العامة"), PS("hl", "Naskh", 11, INK2)),
        dot(status_color(health)),
        Paragraph(ar(health), PS("hw", "Naskh-Bold", 13, status_color(health)))]],
        colWidths=[3.2 * cm, 0.5 * cm, CONTENT_W - 3.7 * cm],
        style=TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                          ("LEFTPADDING", (0, 0), (-1, -1), 0),
                          ("RIGHTPADDING", (0, 0), (-1, -1), 0)])))

    # ── executive summary ──
    if chief.get("executive_summary"):
        story.extend(section("الملخص التنفيذي"))
        story.append(Spacer(1, 6))
        story.append(Table([[Paragraph(ar(chief["executive_summary"],
                                          CONTENT_W - 0.35 * cm - 8, "Naskh", 11),
                                       PS("sum", "Naskh", 11, INK, leading=18)), ""]],
            colWidths=[CONTENT_W - 0.35 * cm, 0.35 * cm],
            style=TableStyle([("BACKGROUND", (1, 0), (1, 0), ACC),
                              ("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("TOPPADDING", (0, 0), (0, -1), 4),
                              ("BOTTOMPADDING", (0, 0), (0, -1), 4),
                              ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (0, -1), 8)])))

    # ── KPIs ──
    kpis = _rows(chief.get("kpis"))
    if kpis:
        story.extend(section("المؤشرات الرئيسية"))
        story.append(Spacer(1, 4))
        data = [[Paragraph(ar(h), PS("th", "Naskh-Bold", 9, INK2, TA_CENTER))
                 for h in ["المؤشر", "القيمة", "الاتجاه", "الحالة"]]]
        for k in kpis:
            k = _obj(k)
            s = k.get("status", "")
            data.append([Paragraph(ar(k.get("name", ""), 8 * cm - 8, "Naskh", 10), PS("kn", "Naskh", 10, INK)),
                         val(k.get("value", "")),
                         val(k.get("trend", ""), color=INK2),
                         val(s, color=status_color(s), bold=True)])
        story.append(Table(data, colWidths=[8 * cm, 3.5 * cm, 2 * cm, 3.5 * cm],
                           style=TableStyle(rows_style(len(data), header=True))))

    # ── top actions ──
    actions = _rows(chief.get("top_actions"))
    if actions:
        story.extend(section("خطة العمل الفورية"))
        story.append(Spacer(1, 4))
        for i, act in enumerate(actions, 1):
            act = _obj(act)
            meta = "  ·  ".join(p for p in (act.get("owner", ""), act.get("deadline", ""),
                                            act.get("impact", "")) if p)
            row = Table([[Paragraph(ar(act.get("action", ""), CONTENT_W - 1.2 * cm - 8, "Naskh-Bold", 10),
                                    PS("aa", "Naskh-Bold", 10, INK))],
                         [Paragraph(ar(meta, CONTENT_W - 1.2 * cm - 8, "Naskh", 8.5),
                                    PS("am", "Naskh", 8.5, INK2))]],
                colWidths=[CONTENT_W],
                style=TableStyle([("LINEBELOW", (0, -1), (-1, -1), 0.4, LINE),
                                  ("TOPPADDING", (0, 0), (-1, 0), 8),
                                  ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                  ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
            story.append(Table([[row, Paragraph(ar(arabic_num(_int(act.get("priority", i), i))),
                                                PS("an", "Naskh-Bold", 11, ACC))]],
                colWidths=[CONTENT_W - 1.2 * cm, 1.2 * cm],
                style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                  ("RIGHTPADDING", (0, 0), (-1, -1), 0)])))

    # ── risk register ──
    risks = _rows(risk.get("risks"))
    if risks:
        story.extend(section("سجل المخاطر"))
        story.append(Spacer(1, 4))
        data = [[Paragraph(ar(h), PS("rh", "Naskh-Bold", 9, INK2, TA_CENTER))
                 for h in ["المخاطرة", "المستوى", "الوصف", "الحل"]]]
        for rk in risks:
            rk = _obj(rk)
            lv = rk.get("level", "")
            data.append([Paragraph(ar(rk.get("title", ""), 4.5 * cm - 8, "Naskh-Bold", 9), PS("rt", "Naskh-Bold", 9, INK)),
                         val(lv, color=status_color(lv), size=9, bold=True),
                         Paragraph(ar(rk.get("description", ""), 5 * cm - 8, "Naskh", 8.5), PS("rd", "Naskh", 8.5, INK2)),
                         Paragraph(ar(rk.get("solution", ""), 5 * cm - 8, "Naskh", 8.5), PS("rs", "Naskh", 8.5, INK2))])
        story.append(Table(data, colWidths=[4.5 * cm, 2.5 * cm, 5 * cm, 5 * cm],
                           style=TableStyle(rows_style(len(data), header=True) +
                                            [("VALIGN", (0, 0), (-1, -1), "TOP")])))

    # ── schedule ──
    story.extend(section("الجدول الزمني"))
    story.append(Spacer(1, 4))
    sdata = [[Paragraph(ar(lbl, 10 * cm - 8, "Naskh", 10), PS("sl", "Naskh", 10, INK2)), val(v)]
             for lbl, v in (("التأخير (يوم)", sched.get("delay_days", "-")),
                            ("الإنهاء الأصلي", sched.get("original_end", "-")),
                            ("الإنهاء الجديد", sched.get("new_end", "-")))]
    story.append(Table(sdata, colWidths=[10 * cm, 7 * cm],
                       style=TableStyle(rows_style(len(sdata)))))

    # ── finance ──
    story.extend(section("الوضع المالي"))
    story.append(Spacer(1, 4))
    dev = _num(fin.get("deviation_pct", 0))
    fdata = [[Paragraph(ar(lbl, 10 * cm - 8, "Naskh", 10), PS("fl", "Naskh", 10, INK2)), val(v)]
             for lbl, v in (("الميزانية الإجمالية", fin.get("total_budget", "-")),
                            ("المُنفَق", fin.get("spent", "-")),
                            ("المتبقي", fin.get("remaining", "-")),
                            ("نسبة الإنفاق", _pct(fin.get("spent_pct", 0))))]
    fdata.append([Paragraph(ar("نسبة الانحراف"), PS("fld", "Naskh", 10, INK2)),
                  val(f"{dev:+g}%", color=BAD if dev > 0 else INK)])
    story.append(Table(fdata, colWidths=[10 * cm, 7 * cm],
                       style=TableStyle(rows_style(len(fdata)))))

    # ── page footer (drawn on every page, never stranded) ──
    def _footer(canv, _doc):
        canv.saveState()
        y = 1.2 * cm
        canv.setStrokeColor(LINE)
        canv.setLineWidth(1)
        canv.line(2 * cm, y, A4[0] - 2 * cm, y)
        canv.setFont("Naskh", 8)
        canv.setFillColor(INK3)
        canv.drawCentredString(A4[0] / 2, y - 0.45 * cm,
            ar(f'مرصد — LTT PMO · {datetime.datetime.now().strftime("%Y-%m-%d %H:%M")} · سري'))
        canv.restoreState()

    doc = SimpleDocTemplate(output_path, pagesize=A4,
        rightMargin=2 * cm, leftMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title="مرصد — تقرير حالة المشروع", author="marsad")
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output_path


def export_excel(results: dict, output_path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils  import get_column_letter

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    wb  = Workbook()
    MID = "1E3A5F"; ACC = "2563EB"
    WHITE = "FFFFFF"; GRAY = "F1F5F9"; ALT = "EBF3FB"
    GRN = "D1FAE5"; RED = "FEE2E2"; YEL = "FEF3C7"
    BORD_C = "CBD5E1"

    # one status→colour table, shared by every ladder below (health banner,
    # KPI status column, risk level, schedule phase status) — see core/status.py
    _XL_TIER      = {"good": "065F46", "warn": "78350F", "bad": "7F1D1D", "neutral": "334155"}
    _XL_TIER_FG   = {"good": "6EE7B7", "warn": "FCD34D", "bad": "FCA5A5", "neutral": BORD_C}
    _XL_TIER_FILL = {"good": GRN,      "warn": YEL,      "bad": RED,      "neutral": GRAY}

    def hdr(cell, bg=MID, fg=WHITE, bold=True, size=11):
        cell.font      = Font(bold=bold, color=fg, size=size, name="Arial")
        cell.fill      = PatternFill("solid", start_color=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def body(cell, bold=False, color="1E293B", bg=None, center=False):
        cell.font      = Font(bold=bold, color=color, size=10, name="Arial")
        cell.alignment = Alignment(horizontal="center" if center else "right",
                                   vertical="center", wrap_text=True)
        if bg: cell.fill = PatternFill("solid", start_color=bg)

    def border_all(ws, r1, r2, c1, c2):
        s = Side(style="thin", color=BORD_C)
        b = Border(left=s, right=s, top=s, bottom=s)
        for row in ws.iter_rows(r1,r2,c1,c2):
            for cell in row: cell.border = b

    chief  = _obj(results.get("chief"))
    risks  = _rows(_obj(results.get("risk")).get("risks"))
    sched  = _obj(results.get("schedule"))
    fin    = _obj(results.get("cost"))
    qual   = _obj(results.get("quality"))
    safety = _obj(results.get("safety"))
    health = chief.get("overall_health","")
    hbg    = _XL_TIER[tier(health)]
    hfg    = _XL_TIER_FG[tier(health)]
    CW     = 9638

    # ═══ ورقة 1: لوحة التحكم ═══
    ws1 = wb.active; ws1.title="لوحة التحكم"; ws1.sheet_view.rightToLeft=True
    ws1.merge_cells("A1:G1"); hdr(ws1["A1"],"0D1B2A",WHITE,True,16)
    ws1["A1"].value = "نظام إدارة المشاريع الذكي — LTT 4G/5G — المنطقة الوسطى"
    ws1.row_dimensions[1].height=40
    ws1.merge_cells("A2:G2"); hdr(ws1["A2"],"0D1B2A","94A3B8",False,10)
    ws1["A2"].value = f"تاريخ: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1.merge_cells("A4:G4"); hdr(ws1["A4"],hbg,hfg,True,14)
    ws1["A4"].value = f"الحالة العامة:  {health}"
    ws1.row_dimensions[4].height=32

    # الملخص
    ws1.merge_cells("A6:G6"); hdr(ws1["A6"],ACC,WHITE,True,12)
    ws1["A6"].value = "الملخص التنفيذي"
    ws1.merge_cells("A7:G9"); body(ws1["A7"],bg=GRAY)
    ws1["A7"].value = _cell(chief.get("executive_summary",""))
    ws1["A7"].alignment = Alignment(horizontal="right",vertical="top",wrap_text=True)
    ws1.row_dimensions[7].height=70

    # KPIs
    r=11
    ws1.merge_cells(f"A{r}:G{r}"); hdr(ws1[f"A{r}"],ACC,WHITE,True,12)
    ws1[f"A{r}"].value="مؤشرات الأداء"
    for ci,h in enumerate(["المؤشر","القيمة","الاتجاه","الحالة"],1):
        hdr(ws1.cell(r+1,ci),MID); ws1.cell(r+1,ci).value=h
    for i,kpi in enumerate(_rows(chief.get("kpis")),r+2):
        kpi=_obj(kpi)
        s=kpi.get("status","")
        t=tier(s)
        bg=_XL_TIER_FILL[t]
        for ci,v in enumerate([kpi.get("name",""),kpi.get("value",""),kpi.get("trend",""),s],1):
            c=ws1.cell(i,ci); c.value=_cell(v)
            body(c,bold=(ci==4),
                 color=_XL_TIER[t] if ci==4 else "1E293B",
                 bg=bg if ci==4 else (GRAY if i%2==0 else WHITE),center=(ci>1))
    for ci,w in enumerate([30,15,10,15,20,15,15],1):
        ws1.column_dimensions[get_column_letter(ci)].width=w

    # خطة العمل
    act_start=r+2+len(_rows(chief.get("kpis")))+2
    ws1.merge_cells(f"A{act_start}:G{act_start}"); hdr(ws1[f"A{act_start}"],MID,WHITE,True,12)
    ws1[f"A{act_start}"].value="خطة العمل الفورية"
    for ci,h in enumerate(["#","الإجراء","المسؤول","الموعد","التأثير"],1):
        hdr(ws1.cell(act_start+1,ci),MID); ws1.cell(act_start+1,ci).value=h
    pbgs=["7F1D1D","78350F","1E3A5F"]
    for i,act in enumerate(_rows(chief.get("top_actions")),act_start+2):
        act=_obj(act)
        pr=_int(act.get("priority",1),1)
        p=max(0,min(pr-1,2)); bg=pbgs[p]
        for ci,v in enumerate([str(pr),act.get("action",""),act.get("owner",""),
                                act.get("deadline",""),act.get("impact","")],1):
            c=ws1.cell(i,ci); c.value=_cell(v)
            body(c,bold=(ci==1),
                 color="FCA5A5" if ci==1 and p==0 else "FCD34D" if ci==1 and p==1 else "93C5FD" if ci==1 else "1E293B",
                 bg=bg if ci==1 else (GRAY if i%2==0 else WHITE),center=(ci==1))
        ws1.row_dimensions[i].height=22

    # ═══ ورقة 2: المخاطر ═══
    ws2=wb.create_sheet("المخاطر"); ws2.sheet_view.rightToLeft=True
    ws2.merge_cells("A1:F1"); hdr(ws2["A1"],"0D1B2A",WHITE,True,13)
    ws2["A1"].value="سجل المخاطر الشامل"; ws2.row_dimensions[1].height=32
    for ci,(h,w) in enumerate(zip(["المخاطرة","المستوى","الفئة","الوصف","الحل","المسؤول"],
                                   [25,14,14,38,38,20]),1):
        hdr(ws2.cell(2,ci),MID); ws2.cell(2,ci).value=h
        ws2.column_dimensions[get_column_letter(ci)].width=w
    for ri,rk in enumerate(risks,3):
        rk=_obj(rk)
        l=rk.get("level",""); t=tier(l); bg,fg=_XL_TIER[t],_XL_TIER_FG[t]
        for ci,v in enumerate([rk.get("title",""),l,rk.get("category",""),
                                rk.get("description",""),rk.get("solution",""),rk.get("owner","")],1):
            c=ws2.cell(ri,ci); c.value=_cell(v)
            body(c,bold=(ci==2),color=fg if ci==2 else "1E293B",
                 bg=bg if ci==2 else (GRAY if ri%2==0 else WHITE),center=(ci==2))
        ws2.row_dimensions[ri].height=28

    # ═══ ورقة 3: الجدول الزمني ═══
    ws3=wb.create_sheet("الجدول الزمني"); ws3.sheet_view.rightToLeft=True
    ws3.merge_cells("A1:D1"); hdr(ws3["A1"],"0D1B2A",WHITE,True,13)
    ws3["A1"].value="الجدول الزمني للمشروع"; ws3.row_dimensions[1].height=32
    for ri,(l,v,vbg) in enumerate([("التأخير (يوم)",sched.get("delay_days","-"),"FEE2E2"),
                                    ("الإنهاء الأصلي",sched.get("original_end","-"),GRAY),
                                    ("الإنهاء الجديد",sched.get("new_end","-"),"FEF3C7")],2):
        body(ws3.cell(ri,1),True,bg=GRAY); ws3.cell(ri,1).value=l
        body(ws3.cell(ri,2),True,bg=vbg,center=True); ws3.cell(ri,2).value=str(v)
        ws3.row_dimensions[ri].height=24
    ws3.column_dimensions["A"].width=28; ws3.column_dimensions["B"].width=22
    r2=6; ws3.merge_cells(f"A{r2}:D{r2}"); hdr(ws3[f"A{r2}"],ACC,WHITE,True,11)
    ws3[f"A{r2}"].value="مراحل المشروع"
    for ci,(h,w) in enumerate(zip(["المرحلة","الحالة","الإنجاز (%)"],
                                   [28,16,18]),1):
        hdr(ws3.cell(r2+1,ci),MID); ws3.cell(r2+1,ci).value=h
        ws3.column_dimensions[get_column_letter(ci)].width=w
    for ri,ph in enumerate(_rows(sched.get("phases")),r2+2):
        ph=_obj(ph)
        s=ph.get("status",""); t=tier(s); sbg=_XL_TIER_FILL[t]
        for ci,v in enumerate([ph.get("name",""),s,_pct(ph.get("completion_pct",0))],1):
            c=ws3.cell(ri,ci); c.value=_cell(v)
            body(c,bold=(ci==2),
                 color=_XL_TIER[t] if ci==2 else "1E293B",
                 bg=sbg if ci==2 else (GRAY if ri%2==0 else WHITE),center=(ci>1))
        ws3.row_dimensions[ri].height=22

    # ═══ ورقة 4: المالية ═══
    ws4=wb.create_sheet("المالية"); ws4.sheet_view.rightToLeft=True
    ws4.merge_cells("A1:C1"); hdr(ws4["A1"],"0D1B2A",WHITE,True,13)
    ws4["A1"].value="الوضع المالي"; ws4.row_dimensions[1].height=32
    for ri,(l,v,bg) in enumerate([
        ("الميزانية الإجمالية",fin.get("total_budget","-"),"D1FAE5"),
        ("المُنفَق",fin.get("spent","-"),"DBEAFE"),
        ("المتبقي",fin.get("remaining","-"),"E0F2FE"),
        ("نسبة الإنفاق",_pct(fin.get("spent_pct",0)),GRAY),
        ("نسبة الانحراف",f"{_num(fin.get('deviation_pct', 0)):+g}%","FEE2E2"),
    ],2):
        ws4.row_dimensions[ri].height=24
        body(ws4.cell(ri,1),True,bg=GRAY); ws4.cell(ri,1).value=l
        body(ws4.cell(ri,2),True,bg=bg,center=True); ws4.cell(ri,2).value=str(v)
    ws4.column_dimensions["A"].width=28; ws4.column_dimensions["B"].width=22

    # ═══ ورقة 5: الجودة والسلامة ═══
    ws5=wb.create_sheet("الجودة والسلامة"); ws5.sheet_view.rightToLeft=True
    ws5.merge_cells("A1:C1"); hdr(ws5["A1"],"0D1B2A",WHITE,True,13)
    ws5["A1"].value="تقرير الجودة والسلامة"; ws5.row_dimensions[1].height=32
    for ri,(l,v,bg) in enumerate([
        ("محطات فُحصت",qual.get("inspected","-"),GRAY),
        ("اجتازت المعيار",qual.get("passed","-"),GRN),
        ("تحتاج مراجعة",qual.get("failed","-"),RED),
        ("نسبة النجاح",_pct(qual.get("pass_rate",0)),GRN),
        ("حوادث السلامة",safety.get("incidents","-"),RED),
        ("درجة السلامة",_pct(safety.get("safety_score",0)),GRN),
    ],2):
        ws5.row_dimensions[ri].height=24
        body(ws5.cell(ri,1),True,bg=GRAY); ws5.cell(ri,1).value=l
        body(ws5.cell(ri,2),True,bg=bg,center=True); ws5.cell(ri,2).value=str(v)
    ws5.column_dimensions["A"].width=28; ws5.column_dimensions["B"].width=20

    # ── ورقة الإدارات — الوكلاء الذين لا تظهر نتائجهم في أي مكان آخر ──
    DEPT_ROWS = (
        ("العمليات الميدانية", "ops", (
            ("نسبة الإنجاز", lambda d: _pct(d.get("completion_pct", 0))),
            ("المواقع النشطة", lambda d: _int(d.get("active_sites", 0))),
            ("حالة الفريق", lambda d: d.get("team_status", "-")))),
        ("الأعمال الإنشائية", "civil", (
            ("الأبراج المنجزة", lambda d: _int(d.get("towers_built", 0))),
            ("إجمالي الأبراج", lambda d: _int(d.get("towers_total", 0))),
            ("نسبة الإنشاء", lambda d: _pct(d.get("civil_pct", 0))))),
        ("العقود", "contract", (
            ("العقود النشطة", lambda d: _int(d.get("active_contracts", 0))),
            ("القيمة الإجمالية", lambda d: d.get("total_value", "-")),
            ("مدفوعات معلّقة", lambda d: d.get("pending_payments", "-")))),
        ("المشتريات", "procure", (
            ("طلبات معلّقة", lambda d: _int(d.get("pending_orders", 0))),
            ("موردون معتمدون", lambda d: _int(d.get("approved_vendors", 0))),
            ("قيمة أوامر الشراء", lambda d: d.get("total_po_value", "-")))),
        ("المخازن وسلاسل التوريد", "supply", (
            ("امتلاء المخزن", lambda d: _pct(d.get("warehouse_fill_pct", 0))),
            ("شحنات في الطريق", lambda d: _int(d.get("in_transit_shipments", 0))),
            ("شحنات متأخرة", lambda d: _int(d.get("delayed_shipments", 0))))),
    )

    wsd = wb.create_sheet("الإدارات")
    wsd.sheet_view.rightToLeft = True
    for ci, h in enumerate(["الإدارة", "المؤشر", "القيمة"], 1):
        hdr(wsd.cell(1, ci), MID)
        wsd.cell(1, ci).value = h
    r = 2
    for label, key, fields in DEPT_ROWS:
        data = _obj(results.get(key))
        if not data or "error" in data:
            continue
        for field_label, getter in fields:
            wsd.cell(r, 1).value = label
            wsd.cell(r, 2).value = field_label
            try:
                wsd.cell(r, 3).value = getter(data)
            except Exception:
                wsd.cell(r, 3).value = "-"
            for ci in (1, 2, 3):
                body(wsd.cell(r, ci), bg=(GRAY if r % 2 == 0 else WHITE))
            r += 1
    wsd.column_dimensions["A"].width = 28
    wsd.column_dimensions["B"].width = 24
    wsd.column_dimensions["C"].width = 18

    wb.save(output_path)
    return output_path
